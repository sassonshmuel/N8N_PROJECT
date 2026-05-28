#!/usr/bin/env python3

import json
import sys
from pathlib import Path
from docx import Document


def extract_docx(docx_path: str) -> dict:
    docx_file = Path(docx_path)

    if not docx_file.exists():
        raise FileNotFoundError(f"DOCX file not found: {docx_path}")

    doc = Document(docx_path)

    paragraphs = []
    tables = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            paragraphs.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        rows = []

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)

        tables.append({
            "table": table_index,
            "rows": rows,
        })

    text_parts = []

    if paragraphs:
        text_parts.append("\n".join(paragraphs))

    for table in tables:
        text_parts.append(f"\n--- Table {table['table']} ---")
        for row in table["rows"]:
            text_parts.append(" | ".join(row))

    full_text = "\n".join(text_parts).strip()

    return {
        "file_name": docx_file.name,
        "file_path": str(docx_file),
        "file_type": "docx",
        "text": full_text,
        "has_text": bool(full_text),
        "has_images": False,
        "images": [],
        "paragraphs": paragraphs,
        "tables": tables,
        "total_paragraphs": len(paragraphs),
        "total_tables": len(tables),
    }


def main():
    if len(sys.argv) < 2:
        print(json.dumps({
            "success": False,
            "error": "Usage: python3 extract_docx.py <docx_path>",
        }))
        sys.exit(1)

    docx_path = sys.argv[1]

    try:
        result = extract_docx(docx_path)
        result["success"] = True
        print(json.dumps(result, ensure_ascii=False))

    except Exception as e:
        print(json.dumps({
            "success": False,
            "error": str(e),
            "file_path": docx_path,
        }))
        sys.exit(1)


if __name__ == "__main__":
    main()
